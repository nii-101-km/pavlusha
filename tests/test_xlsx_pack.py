"""XLSX capability boundary tests; optional document dependencies are explicit."""
import base64
import importlib.util
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stdout, redirect_stderr
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from pavlusha_agent.functions import FunctionRegistry

PACK = Path(__file__).resolve().parents[1] / 'tools/xlsx_functions.py'
DOCX = PACK.with_name('docx_functions.py')
_HAVE_XL = bool(importlib.util.find_spec('openpyxl') and importlib.util.find_spec('lxml'))
if _HAVE_XL:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.formula import ArrayFormula


def load(root, argv=None):
    with patch.object(sys, 'argv', argv or ['agent.py', '--workdir', str(root)]):
        spec = importlib.util.spec_from_file_location('xlsx_test_pack', PACK)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
    return module


def fixture(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Plan'
    ws.append(['Item', 'Count', 'Rate', 'Amount'])
    ws.append(['A', 2, 100, '=B2*C2'])
    ws.append(['B', 3, 200, '=B3*C3'])
    ws.append(['Total', None, None, '=SUM(D2:D3)'])
    ws['A6'] = 'Merged title'
    ws.merge_cells('A6:D6')
    ws['A7'] = 'literal =SUM(A1:A2)'
    ws['C2'].number_format = '#,##0.00'
    ws['C2'].font = Font(bold=True, color='112233')
    ws['A1'].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 27
    ws.row_dimensions[7].hidden = True
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].hidden = True
    ws.freeze_panes = 'B2'
    ws.print_area = 'A1:D6'
    ws.print_title_rows = '1:1'
    ws.page_setup.orientation = 'landscape'
    ws.auto_filter.ref = 'A1:D3'
    dv = DataValidation(type='whole', operator='greaterThan', formula1=0)
    dv.add('B2:B3')
    ws.add_data_validation(dv)
    wb.defined_names.add(DefinedName('Total', attr_text="'Plan'!$D$4"))
    hidden = wb.create_sheet('History')
    hidden.sheet_state = 'veryHidden'
    hidden['A1'] = '=Plan!D4'
    hidden['A2'] = 999
    wb.save(path)
    # Known deliberately stored caches, not computed by openpyxl.
    rewrite(path, {'xl/worksheets/sheet1.xml': lambda b: b.replace(b'<f>B2*C2</f><v></v>', b'<f>B2*C2</f><v>200</v>'),
                   'xl/worksheets/sheet2.xml': lambda b: b.replace(b'<f>Plan!D4</f><v></v>', b'<f>Plan!D4</f><v>800</v>')})


def rewrite(path, replacements):
    stream = io.BytesIO()
    with ZipFile(path) as old, ZipFile(stream, 'w') as new:
        for info in old.infolist():
            data = old.read(info.filename)
            if info.filename in replacements:
                data = replacements[info.filename](data)
            new.writestr(info, data)
    path.write_bytes(stream.getvalue())


class XlsxRegistryTests(unittest.TestCase):
    def test_repeatable_docx_xlsx_both_orders_no_new_public_helpers(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(sys, 'argv', ['agent.py', '--workdir', tmp]):
            for modules in ([DOCX, PACK], [PACK, DOCX]):
                registry = FunctionRegistry(modules)
                expected = [prefix + '_' + suffix for prefix in (['docx', 'xlsx'] if modules[0] == DOCX else ['xlsx', 'docx'])
                            for suffix in ('inspect', 'read', 'edit', 'render')]
                self.assertEqual(list(registry.functions), expected)
                self.assertEqual([d['name'] for d in registry.descriptions], expected)

    def test_root_is_existing_cli_workdir_frozen_and_ignores_docx_root(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, {'PAVLUSHA_DOCX_ROOT': '/unused'}):
            root = Path(tmp)
            for options in (['--workdir', tmp], ['--workdir=' + tmp], ['--workdir', '/unused', '--workdir', tmp]):
                pack = load(root, ['agent.py', *options])
                with patch('os.getcwd', return_value='/tmp'), patch.object(sys, 'argv', ['agent.py', '--workdir', '/changed']):
                    self.assertEqual(pack._path('file.xlsx'), root / 'file.xlsx')
            with patch('os.getcwd', return_value=tmp):
                self.assertEqual(load(root, ['agent.py'])._ROOT, root / 'agent-work')

    def test_disabled_runtime_does_not_load_packs_and_keeps_shell_flow(self):
        from tests.test_runtime_lifecycle import run_script
        from tests.test_reasoning_window import init_turn, turn
        from tests.test_checkpoint_snapshots import done
        with tempfile.TemporaryDirectory() as tmp, patch.object(FunctionRegistry, '_load_module') as loader:
            seen, result, error = run_script(Path(tmp), [init_turn(),
                turn({'action': 'shell', 'command': 'verify'}), done('verify'),
                turn({'action': 'finish', 'summary': 'done'})])
            self.assertEqual(result, 0, error)
            self.assertNotIn('xlsx_inspect', str(seen))
            self.assertNotIn('User-enabled external capabilities', str(seen))
            loader.assert_not_called()


@unittest.skipUnless(_HAVE_XL, 'requires optional openpyxl/lxml; run with document Python')
class XlsxPackTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.pack = load(self.root)
        fixture(self.root / 'plan.xlsx')
        self.info = self.pack.xlsx_inspect('plan.xlsx')
        self.assertTrue(self.info['ok'], self.info)
        self.rev = self.info['revision']

    def error(self, result, code):
        self.assertFalse(result['ok'], result)
        self.assertEqual(result['error']['code'], code, result)

    def read(self, cell_range='A1:D7', path='plan.xlsx', revision=None, sheet='Plan', **kwargs):
        records, cursor = [], None
        for _ in range(100):
            result = self.pack.xlsx_read(path, revision or self.rev, sheet, cell_range, cursor=cursor, **kwargs)
            self.assertTrue(result['ok'], result)
            self.assertLessEqual(len(self.pack._json(result)), kwargs.get('max_chars', 8000))
            records.extend(result['items'])
            cursor = result['next_cursor']
            if cursor is None:
                return {r['cell']: r for r in records}
        self.fail('cursor failed to terminate')

    def edit(self, edits=None, output='copy.xlsx', revision=None):
        return self.pack.xlsx_edit('plan.xlsx', output, revision or self.rev,
                                   edits if edits is not None else [{'sheet': 'Plan', 'cell': 'C2', 'value': 150}])

    def test_inspect_pagination_layout_features_and_hidden_sheets(self):
        sheets = [r for r in self.info['items'] if r['kind'] == 'sheet']
        self.assertEqual([r['name'] for r in sheets], ['Plan', 'History'])
        self.assertEqual(sheets[1]['state'], 'veryHidden')
        self.assertEqual(sheets[0]['formula_cells'], 3)
        records, cursor = [], None
        while True:
            result = self.pack.xlsx_inspect('plan.xlsx', 'Plan', cursor, limit=2)
            self.assertTrue(result['ok'], result)
            records.extend(result['items'])
            cursor = result['next_cursor']
            if cursor is None:
                break
        self.assertTrue(any(r['kind'] == 'merge' and r['range'] == 'A6:D6' for r in records))
        self.assertTrue(any(r['kind'] == 'row' and r['hidden'] for r in records))
        self.assertTrue(any(r['kind'] == 'column' and r['hidden'] for r in records))
        self.assertTrue(any(r['kind'] == 'validation' for r in records))
        self.assertTrue(any(r['kind'] == 'defined_name' for r in records))
        layout = next(r for r in records if r['kind'] == 'sheet_layout')
        self.assertIn('$A$1', layout['print_area'])
        self.assertEqual(layout['freeze_panes'], 'B2')

    def test_read_formula_cache_number_format_merges_and_empty_cells(self):
        cells = self.read(max_chars=1800)
        self.assertEqual(cells['D2']['formula'], '=B2*C2')
        self.assertEqual(cells['D2']['value'], 200)
        self.assertIsNone(cells['D3']['value'])
        self.assertEqual(cells['C2']['number_format'], '#,##0.00')
        self.assertTrue(cells['C2']['bold'])
        self.assertEqual(cells['B6']['merge_anchor'], 'A6')
        self.assertIsNone(cells['B6']['value'])
        self.assertTrue(cells['B2']['column_hidden'])
        self.assertTrue(cells['A7']['row_hidden'])
        self.assertEqual(cells['A1']['row_height'], 27)
        self.assertEqual(self.read('A2', sheet='History')['A2']['value'], 999)

    def test_revision_and_selector_bound_cursors(self):
        first = self.pack.xlsx_read('plan.xlsx', self.rev, 'Plan', 'A1:D7', max_chars=1800)
        self.assertIsNotNone(first['next_cursor'])
        self.error(self.pack.xlsx_read('plan.xlsx', self.rev, 'History', 'A1:D7', first['next_cursor']), 'invalid_cursor')
        self.error(self.pack.xlsx_read('plan.xlsx', self.rev, 'Plan', 'A1:D6', first['next_cursor']), 'invalid_cursor')
        self.error(self.pack.xlsx_inspect('plan.xlsx', cursor='bad'), 'invalid_cursor')
        self.error(self.pack.xlsx_read('plan.xlsx', 'old', 'Plan', 'A1'), 'stale_revision')
        self.error(self.edit(revision='old'), 'stale_revision')
        self.error(self.pack.xlsx_render('plan.xlsx', 'old', 'pages'), 'stale_revision')

    def test_value_formula_blank_date_boolean_literal_and_new_cell(self):
        edits = [{'sheet': 'Plan', 'cell': 'C2', 'value': 150},
                 {'sheet': 'Plan', 'cell': 'D4', 'formula': '=SUM(D2:D3)+10'},
                 {'sheet': 'Plan', 'cell': 'A7', 'value': '=literal'},
                 {'sheet': 'Plan', 'cell': 'B3', 'value': None},
                 {'sheet': 'Plan', 'cell': 'F9', 'date': '2026-10-06'},
                 {'sheet': 'Plan', 'cell': 'G9', 'value': True}]
        result = self.edit(edits)
        self.assertTrue(result['ok'], result)
        cells = self.read('A1:G9', 'copy.xlsx', result['revision'])
        self.assertEqual(cells['C2']['value'], 150)
        self.assertEqual(cells['D4']['formula'], '=SUM(D2:D3)+10')
        self.assertIsNone(cells['D2']['value'])
        self.assertIsNone(self.read('A1', 'copy.xlsx', result['revision'], sheet='History')['A1']['value'])
        self.assertEqual(cells['A7']['value'], '=literal')
        self.assertIsNone(cells['A7']['formula'])
        self.assertEqual(cells['F9']['value']['iso'], '2026-10-06')
        self.assertTrue(cells['G9']['value'])
        self.assertTrue(cells['C2']['bold'])
        self.assertEqual(cells['C2']['number_format'], '#,##0.00')

    def test_exact_untouched_zip_payloads_chart_and_relationship_preservation(self):
        wb = openpyxl.load_workbook(self.root / 'plan.xlsx')
        chart = BarChart()
        chart.add_data(Reference(wb['Plan'], min_col=3, min_row=1, max_row=3), titles_from_data=True)
        wb['Plan'].add_chart(chart, 'F1')
        wb.save(self.root / 'plan.xlsx')
        original = (self.root / 'plan.xlsx').read_bytes()
        self.rev = self.pack.xlsx_inspect('plan.xlsx')['revision']
        result = self.edit()
        self.assertTrue(result['ok'], result)
        self.assertEqual((self.root / 'plan.xlsx').read_bytes(), original)
        with ZipFile(self.root / 'plan.xlsx') as old, ZipFile(self.root / 'copy.xlsx') as new:
            self.assertEqual(old.namelist(), new.namelist())
            for name in old.namelist():
                if name not in result['changed_parts']:
                    self.assertEqual(old.read(name), new.read(name), name)
            self.assertEqual(old.read('xl/charts/chart1.xml'), new.read('xl/charts/chart1.xml'))
            self.assertEqual(old.read('xl/styles.xml'), new.read('xl/styles.xml'))

    def test_failed_batch_no_output_and_invalid_values(self):
        for bad in ({'sheet': 'Missing', 'cell': 'A1', 'value': 1},
                    {'sheet': 'Plan', 'cell': 'B6', 'value': 'merge follower'},
                    {'sheet': 'Plan', 'cell': 'A1', 'value': '\x00'},
                    {'sheet': 'Plan', 'cell': 'A1', 'value': []},
                    {'sheet': 'Plan', 'cell': 'A1', 'value': float('nan')},
                    {'sheet': 'Plan', 'cell': 'A1', 'value': 10**400},
                    {'sheet': 'Plan', 'cell': 'A1', 'date': 'not a date'},
                    {'sheet': 'Plan', 'cell': 'A1', 'formula': 'SUM(B1:B2)'},
                    {'sheet': 'Plan', 'cell': 'A1', 'formula': "='[outside.xlsx]S'!A1"},
                    {'sheet': 'Plan', 'cell': 'A1', 'value': 1, 'formula': '=1'},
                    {'sheet': 'Plan', 'cell': 'A1:B2', 'value': 1}):
            with self.subTest(bad=bad):
                result = self.edit([{'sheet': 'Plan', 'cell': 'C2', 'value': 150}, bad])
                self.assertFalse(result['ok'], result)
                self.assertFalse((self.root / 'copy.xlsx').exists())
        self.error(self.edit([{'sheet': 'Plan', 'cell': 'C2', 'value': 1}] * 2), 'invalid_edit')
        self.assertEqual(list(self.root.glob('.xlsx-edit-*')), [])

    def test_stale_after_actual_source_mutation(self):
        rewrite(self.root / 'plan.xlsx', {'xl/worksheets/sheet1.xml': lambda b: b.replace(b'Merged title', b'Changed title')})
        self.error(self.edit(), 'stale_revision')
        self.error(self.pack.xlsx_read('plan.xlsx', self.rev, 'Plan', 'A1'), 'stale_revision')

    def test_cli_repeatable_packs_no_root_env_real_dispatch(self):
        from pavlusha_agent.cli import main
        from tests.test_reasoning_window import init_turn, turn
        from tests.test_checkpoint_snapshots import done
        replies = iter([init_turn(),
            turn({'action': 'call_function', 'name': 'xlsx_inspect', 'arguments': {'path': 'plan.xlsx'}}),
            turn({'action': 'call_function', 'name': 'xlsx_edit', 'arguments': {
                'path': 'plan.xlsx', 'revision': self.rev, 'output_path': 'copy.xlsx',
                'edits': [{'sheet': 'Plan', 'cell': 'C2', 'value': 150}]}}),
            done('XLSX edits'), turn({'action': 'finish', 'summary': 'done'})])
        seen = []
        def worker(provider, messages, **kwargs):
            seen.append(str(messages))
            return next(replies)
        argv = ['agent.py', '--no-interactive', '--no-live', '--no-network', '--project-map', 'off',
                '--workdir', str(self.root), '--functions', str(DOCX), '--functions', str(PACK),
                '--model', 'scripted', '--worker-context-budget', '40000', '--max-tokens', '1024',
                '--max-steps', '6', '--project-review-every', '0', 'Verify XLSX']
        self.addCleanup(shutil.rmtree, str(self.root) + '.pavlusha-state', True)
        with patch.object(sys, 'argv', argv), patch('pavlusha_agent.provider.ChatProvider.worker_completion', worker), \
             patch('pavlusha_agent.runtime.run_shell', side_effect=AssertionError('unexpected shell')), \
             redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            self.assertEqual(main(), 0)
        self.assertIn('docx_inspect', seen[0])
        self.assertIn('xlsx_inspect', seen[0])
        self.assertIn('invalidated_workbook_wide', str(seen))
        self.assertEqual(openpyxl.load_workbook(self.root / 'copy.xlsx')['Plan']['C2'].value, 150)

    def test_path_escape_inputs_and_all_outputs(self):
        with tempfile.TemporaryDirectory() as outside:
            (self.root / 'escape').symlink_to(outside, target_is_directory=True)
            for path in ('../x.xlsx', str(self.root / 'plan.xlsx'), 'escape/x.xlsx'):
                self.error(self.pack.xlsx_inspect(path), 'path_escape')
                self.error(self.pack.xlsx_read(path, self.rev, 'Plan', 'A1'), 'path_escape')
                self.error(self.edit(output=path), 'path_escape')
                self.error(self.pack.xlsx_render('plan.xlsx', self.rev, path), 'path_escape')
            self.assertEqual(list(Path(outside).iterdir()), [])

    def test_no_overwrite_publication_failure_and_receipt_bound(self):
        (self.root / 'copy.xlsx').write_bytes(b'keep')
        self.error(self.edit(), 'destination_exists')
        self.assertEqual((self.root / 'copy.xlsx').read_bytes(), b'keep')
        with patch.object(self.pack.os, 'link', side_effect=OSError('failed')):
            self.error(self.edit(output='failed.xlsx'), 'io_error')
        self.assertFalse((self.root / 'failed.xlsx').exists())
        self.assertEqual(list(self.root.glob('.xlsx-edit-*')), [])
        with patch.object(self.pack, '_BUDGET', 100):
            self.error(self.edit(output='small.xlsx'), 'result_limit_exceeded')
        self.assertFalse((self.root / 'small.xlsx').exists())

    def test_invalid_package_xml_macros_and_limits(self):
        (self.root / 'bad.xlsx').write_bytes(b'bad')
        self.error(self.pack.xlsx_inspect('bad.xlsx'), 'invalid_xlsx')
        original = (self.root / 'plan.xlsx').read_bytes()
        for data in (b'<bad>', b'<!DOCTYPE x><x/>', '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x><x/>'.encode('utf-16')):
            (self.root / 'bad.xlsx').write_bytes(original)
            rewrite(self.root / 'bad.xlsx', {'xl/workbook.xml': lambda b: data})
            self.error(self.pack.xlsx_inspect('bad.xlsx'), 'invalid_xlsx')
        (self.root / 'bad.xlsx').write_bytes(original)
        rewrite(self.root / 'bad.xlsx', {'[Content_Types].xml': lambda b: b.replace(b'spreadsheetml.sheet.main', b'sheet.macroEnabled.main')})
        self.error(self.pack.xlsx_inspect('bad.xlsx'), 'unsupported_workbook')
        (self.root / 'bad.xlsx').write_bytes(original)
        rewrite(self.root / 'bad.xlsx', {'xl/worksheets/sheet1.xml': lambda b: b.replace(b'A6:D6', b'A1:XFD1048576')})
        self.error(self.pack.xlsx_inspect('bad.xlsx'), 'workbook_limit_exceeded')

    def test_protected_special_array_and_external_workbook_refused(self):
        wb = openpyxl.load_workbook(self.root / 'plan.xlsx')
        wb['Plan'].protection.sheet = True
        wb.save(self.root / 'plan.xlsx')
        self.rev = self.pack.xlsx_inspect('plan.xlsx')['revision']
        self.error(self.edit(), 'unsupported_target')
        wb['Plan'].protection.sheet = False
        wb['Plan']['D2'] = ArrayFormula(ref='D2:D3', text='=B2:B3*C2:C3')
        wb.save(self.root / 'plan.xlsx')
        self.rev = self.pack.xlsx_inspect('plan.xlsx')['revision']
        self.assertEqual(self.read('D2')['D2']['formula_kind'], 'ArrayFormula')
        self.error(self.edit(), 'unsupported_workbook')

    def test_external_connections_read_flag_but_edit_render_refused(self):
        with ZipFile(self.root / 'plan.xlsx', 'a') as archive:
            archive.writestr('xl/connections.xml', '<connections xmlns="' + self.pack._NS + '"/>')
        info = self.pack.xlsx_inspect('plan.xlsx')
        self.assertTrue(info['ok'], info)
        self.assertTrue(info['items'][0]['features']['connections'])
        self.rev = info['revision']
        self.error(self.edit(), 'unsupported_workbook')
        self.error(self.pack.xlsx_render('plan.xlsx', self.rev, 'pages'), 'unsupported_workbook')

    def test_calc_chain_removed_with_relationship_and_content_type(self):
        with ZipFile(self.root / 'plan.xlsx', 'a') as archive:
            archive.writestr('xl/calcChain.xml', '<calcChain xmlns="' + self.pack._NS + '"/>')
        rewrite(self.root / 'plan.xlsx', {
            'xl/_rels/workbook.xml.rels': lambda b: b.replace(b'</Relationships>',
                b'<Relationship Id="rIdChain" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>'),
            '[Content_Types].xml': lambda b: b.replace(b'</Types>',
                b'<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>')})
        self.rev = self.pack.xlsx_inspect('plan.xlsx')['revision']
        result = self.edit()
        self.assertTrue(result['ok'], result)
        self.assertEqual(result['removed_parts'], ['xl/calcChain.xml'])
        with ZipFile(self.root / 'copy.xlsx') as archive:
            self.assertNotIn('xl/calcChain.xml', archive.namelist())
            self.assertNotIn(b'calcChain', archive.read('[Content_Types].xml'))
            self.assertNotIn(b'calcChain', archive.read('xl/_rels/workbook.xml.rels'))

    def test_shared_formula_target_refused_but_unrelated_edit_preserves_it(self):
        rewrite(self.root / 'plan.xlsx', {'xl/worksheets/sheet1.xml': lambda b:
            b.replace(b'<f>B2*C2</f>', b'<f t="shared" si="0" ref="D2:D3">B2*C2</f>')})
        self.rev = self.pack.xlsx_inspect('plan.xlsx')['revision']
        self.error(self.edit([{'sheet': 'Plan', 'cell': 'D2', 'value': 0}]), 'unsupported_target')
        result = self.edit()
        self.assertTrue(result['ok'], result)
        with ZipFile(self.root / 'copy.xlsx') as archive:
            self.assertIn(b't="shared"', archive.read('xl/worksheets/sheet1.xml'))

    def test_range_unknown_sheet_limits_and_long_cell_are_explicit(self):
        for cell_range in ('A0', 'A1:A', 'A1:D1048576', 'D4:A1', 'XFE1', '$A$1', 'Plan!A1'):
            result = self.pack.xlsx_read('plan.xlsx', self.rev, 'Plan', cell_range)
            self.assertFalse(result['ok'], result)
        self.error(self.pack.xlsx_read('plan.xlsx', self.rev, 'Missing', 'A1'), 'unknown_sheet')
        self.error(self.pack.xlsx_read('plan.xlsx', self.rev, 'Plan', 'A1', max_chars=1), 'invalid_limit')
        result = self.edit([{'sheet': 'Plan', 'cell': 'A7', 'value': 'x' * 30000}])
        cells = self.read('A7', 'copy.xlsx', result['revision'])
        self.assertTrue(cells['A7']['value']['truncated'])
        self.assertEqual(cells['A7']['value']['length'], 30000)

    def test_render_missing_failure_timeout_cleanup_and_no_clobber(self):
        with patch.object(self.pack, '_SOFFICE', None):
            self.error(self.pack.xlsx_render('plan.xlsx', self.rev, 'pages'), 'dependency_missing')
        executable = self.root / 'fake-office'
        executable.write_text('#!/bin/sh\nexit 2\n')
        executable.chmod(0o755)
        with patch.object(self.pack, '_SOFFICE', str(executable)):
            # Poppler is separately required; configure only the existing PATH.
            with patch.object(self.pack.shutil, 'which', return_value='/fake/poppler'):
                self.error(self.pack.xlsx_render('plan.xlsx', self.rev, 'pages'), 'render_failed')
                executable.write_text('#!/bin/sh\nsleep 5\n')
                self.error(self.pack.xlsx_render('plan.xlsx', self.rev, 'pages', timeout_seconds=1), 'render_timeout')
        self.assertFalse((self.root / 'pages').exists())
        self.assertEqual(list(self.root.glob('.xlsx-render-*')), [])
        stage, output = self.root / 'stage', self.root / 'existing'
        stage.mkdir()
        output.mkdir()
        with self.assertRaises(self.pack._Error) as caught:
            self.pack._publish(stage, output)
        self.assertEqual(caught.exception.code, 'destination_exists')
        self.assertTrue(stage.is_dir())

    @unittest.skipUnless(os.getenv('PAVLUSHA_TEST_REAL_XLSX_RENDER') == '1', 'opt-in real LibreOffice/Poppler')
    def test_real_render_recalculation_manifest_source_and_hidden_scope(self):
        original = (self.root / 'plan.xlsx').read_bytes()
        result = self.edit()
        rendered = self.pack.xlsx_render('copy.xlsx', result['revision'], 'pages', dpi=72)
        self.assertTrue(rendered['ok'], rendered)
        cells = self.read('D2:D4', rendered['recalculated_path'], rendered['recalculated_revision'])
        self.assertEqual(cells['D2']['value'], 300)
        self.assertEqual(cells['D4']['value'], 900)
        self.assertEqual((self.root / 'plan.xlsx').read_bytes(), original)
        self.assertEqual(rendered['visual_verification'], 'not_performed')
        manifest = json.loads((self.root / rendered['manifest_path']).read_text())
        self.assertEqual(manifest['recalculated_revision'], rendered['recalculated_revision'])
        self.assertEqual(len(manifest['pages']), rendered['page_count'])
        self.assertEqual(rendered['page_count'], 1)  # Hidden History tab is not silently printed.
        self.error(self.pack.xlsx_render('copy.xlsx', result['revision'], 'pages'), 'destination_exists')


if __name__ == '__main__':
    unittest.main()
