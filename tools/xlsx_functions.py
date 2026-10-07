"""Trusted, stateless XLSX capabilities rooted in the existing CLI --workdir.

openpyxl reads values/styles; edits patch OOXML without an openpyxl round trip.
LibreOffice and Poppler on PATH provide optional recalculation/print rendering.
"""
from __future__ import annotations

import base64
import copy
import ctypes
import datetime as _dt
import errno
import hashlib
import io
import json
import math
import os
import posixpath
import re
import shutil
import signal
import subprocess
import tempfile
import time
import warnings
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from pavlusha_agent.cli import build_parser as _build_parser

try:
    import openpyxl as _xl
    from lxml import etree as _ET
except ImportError:
    _xl = _ET = None

_ROOT = Path(_build_parser().parse_args().workdir).expanduser().resolve()
_SOFFICE = shutil.which('soffice')
_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
_BUDGET = 12000
_CACHE_NOTE = 'Stored formula caches may be absent or stale; openpyxl does not calculate formulas.'


class _Error(Exception):
    def __init__(self, code, message):
        self.code = code
        super().__init__(message)


def _fail(code, message):
    raise _Error(code, message)


def _json(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(',', ':'))


def _call(function, *args):
    try:
        result = function(*args)
        if len(_json(result)) > _BUDGET:
            _fail('result_limit_exceeded', 'Request a smaller batch.')
        return result
    except _Error as exc:
        return {'ok': False, 'error': {'code': exc.code, 'message': str(exc)[:1000]}}
    except (OSError, BadZipFile) as exc:
        return {'ok': False, 'error': {'code': 'io_error', 'message': str(exc)[:1000]}}


def _path(value):
    if not _ROOT.is_dir():
        _fail('configuration_error', 'Runtime workdir does not exist.')
    if type(value) is not str or not value or len(value) > 512 or '\x00' in value:
        _fail('invalid_path', 'Use a nonempty relative path of at most 512 characters.')
    path = Path(value)
    if path.is_absolute() or '..' in path.parts:
        _fail('path_escape', 'Only paths relative to runtime workdir are accepted.')
    try:
        resolved = (_ROOT / path).resolve()
    except (OSError, RuntimeError, ValueError):
        _fail('invalid_path', 'Path cannot be resolved.')
    if not resolved.is_relative_to(_ROOT) or resolved == _ROOT:
        _fail('path_escape', 'Path resolves outside workdir or selects its root.')
    return resolved


def _destination(value):
    output = _path(value)
    if output.exists() or (_ROOT / value).is_symlink():
        _fail('destination_exists', 'Choose a new output path; overwriting is refused.')
    if not output.parent.is_dir():
        _fail('invalid_path', 'Output parent directory must already exist.')
    return output


def _tag(name):
    return '{' + _NS + '}' + name


def _xml(data):
    if b'<!DOCTYPE' in data.upper():
        _fail('invalid_xlsx', 'DTD declarations are refused.')
    try:
        root = _ET.fromstring(data, _ET.XMLParser(resolve_entities=False, no_network=True))
        if root.getroottree().docinfo.doctype:
            _fail('invalid_xlsx', 'DTD declarations are refused.')
        return root
    except _ET.XMLSyntaxError:
        _fail('invalid_xlsx', 'Malformed package XML.')


def _range(value):
    if type(value) is not str or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}(:[A-Z]{1,3}[1-9][0-9]{0,6})?', value):
        _fail('invalid_range', 'Use an A1 cell or bounded A1:C20 rectangle, without sheet names or $.')
    a, b, c, d = _xl.utils.range_boundaries(value)
    if a > c or b > d or c > 16384 or d > 1048576:
        _fail('invalid_range', 'Range is reversed or outside Excel cell bounds.')
    return a, b, c, d


def _open(path, revision=None):
    source = _path(path)
    if source.suffix.lower() != '.xlsx':
        _fail('unsupported_workbook', 'Only ordinary .xlsx files are supported; no XLS/VBA/macros.')
    if not source.is_file():
        _fail('file_not_found', 'Workbook does not exist.')
    with source.open('rb') as stream:
        data = stream.read(32 * 1024 * 1024 + 1)
    if len(data) > 32 * 1024 * 1024:
        _fail('workbook_limit_exceeded', 'Workbook exceeds 32 MiB.')
    digest = hashlib.sha256(data).hexdigest()
    if revision is not None and revision != digest:
        _fail('stale_revision', 'Workbook changed; inspect/read its current revision.')
    if _xl is None or _ET is None:
        _fail('dependency_missing', 'openpyxl and lxml are required; nothing is installed automatically.')
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if (len(infos) > 2048 or len({i.filename for i in infos}) != len(infos)
                    or sum(i.file_size for i in infos) > 128 * 1024 * 1024):
                _fail('workbook_limit_exceeded', 'Too many/duplicate ZIP members or expanded size above 128 MiB.')
            if any(i.flag_bits & 1 for i in infos):
                _fail('unsupported_workbook', 'Encrypted ZIP is not supported.')
            members = {i.filename: archive.read(i) for i in infos}
            comment = archive.comment
        required = {'[Content_Types].xml', '_rels/.rels', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels'}
        if not required <= members.keys():
            _fail('invalid_xlsx', 'Missing standard XLSX package parts.')
        # Validate XML before passing it to openpyxl's own reader.
        roots = {n: _xml(v) for n, v in members.items() if n.endswith(('.xml', '.rels'))}
        content_types = members['[Content_Types].xml'].lower()
        if (b'macroenabled' in content_types or b'macrosheet' in content_types
                or any('vbaproject' in n.lower() for n in members)):
            _fail('unsupported_workbook', 'Macro-enabled packages are outside this pack.')
        main = roots['xl/workbook.xml']
        if main.tag != _tag('workbook'):
            _fail('unsupported_workbook', 'Only conventional SpreadsheetML is supported.')
        rels = {e.get('Id'): e for e in roots['xl/_rels/workbook.xml.rels']}
        parts = {}
        cells = merged_area = 0
        for sheet in main.findall(_tag('sheets') + '/' + _tag('sheet')):
            rel = rels.get(sheet.get('{' + _REL + '}id'))
            if rel is None or rel.get('TargetMode') == 'External' or not rel.get('Type', '').endswith('/worksheet'):
                _fail('unsupported_workbook', 'Only ordinary worksheet tabs are supported.')
            target = rel.get('Target', '')
            name = posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/' + target)
            if not name.startswith('xl/') or name not in roots or roots[name].tag != _tag('worksheet'):
                _fail('invalid_xlsx', 'Invalid worksheet relationship.')
            if sheet.get('name') in parts or name in parts.values():
                _fail('invalid_xlsx', 'Duplicate sheet names or worksheet targets.')
            parts[sheet.get('name')] = name
            root = roots[name]
            cells += len(root.findall('.//' + _tag('c')))
            for merge in root.findall(_tag('mergeCells') + '/' + _tag('mergeCell')):
                a, b, c, d = _range(merge.get('ref'))
                merged_area += (c - a + 1) * (d - b + 1)
        if not 1 <= len(parts) <= 128 or cells + merged_area > 200000:
            _fail('workbook_limit_exceeded', 'Limit: 128 sheets and 200000 stored cells plus merged-cell area.')
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            workbook = _xl.load_workbook(io.BytesIO(data), data_only=False, keep_links=True)
            cached = _xl.load_workbook(io.BytesIO(data), data_only=True, keep_links=True)
        notices = sorted({str(w.message)[:300] for w in caught})[:5]
    except (BadZipFile, KeyError, ValueError, TypeError, IndexError, EOFError) as exc:
        _fail('invalid_xlsx', 'Cannot read workbook: ' + str(exc)[:300])
    features = {label: sum(n.startswith(prefix) for n in members) for label, prefix in (
        ('drawings', 'xl/drawings/'), ('charts', 'xl/charts/'), ('pivot_parts', 'xl/pivot'),
        ('external_links', 'xl/externalLinks/'), ('query_tables', 'xl/queryTables/'),
        ('comments', 'xl/comments'), ('signatures', '_xmlsignatures/'))}
    features['connections'] = 'xl/connections.xml' in members
    features['extensions'] = sum(len(r.findall('.//' + _tag('extLst'))) for r in roots.values())
    return dict(data=data, revision=digest, members=members, infos=infos, comment=comment,
                roots=roots, parts=parts, workbook=workbook, cached=cached, features=features, notices=notices)


def _sheet(doc, name):
    if name not in doc['parts']:
        _fail('unknown_sheet', 'Sheet does not exist: ' + str(name)[:100])
    return doc['workbook'][name]


def _scalar(value):
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time, _dt.timedelta)):
        return {'kind': type(value).__name__, 'iso': value.isoformat()} if not isinstance(value, _dt.timedelta) else {'kind': 'duration', 'seconds': value.total_seconds()}
    if type(value) is float and not math.isfinite(value):
        _fail('invalid_xlsx', 'Non-finite cell value.')
    if value is None or type(value) in (str, int, float, bool):
        return value
    return {'kind': 'unsupported_value', 'type': type(value).__name__}


def _clip(value, length=1000):
    if type(value) is str and len(value) > length:
        return {'text': value[:length], 'truncated': True, 'length': len(value)}
    return value


def _cursor(token, revision, query, count):
    if token is None:
        return 0
    try:
        if len(token) > 1024:
            raise ValueError()
        rev, q, index = json.loads(base64.urlsafe_b64decode(token))
        if rev != revision or q != query or type(index) is not int or not 0 <= index <= count:
            raise ValueError()
        return index
    except (ValueError, TypeError, UnicodeError):
        _fail('invalid_cursor', 'Cursor belongs to a different revision/selector.')


def _page(doc, records, selector, cursor, budget, limit):
    query = hashlib.sha256(_json(selector).encode()).hexdigest()
    index = _cursor(cursor, doc['revision'], query, len(records))
    result = {'ok': True, 'revision': doc['revision'], 'cache_notice': _CACHE_NOTE, 'items': [], 'next_cursor': None}
    while index < len(records) and len(result['items']) < limit:
        token = base64.urlsafe_b64encode(_json([doc['revision'], query, index + 1]).encode()).decode()
        result['items'].append(records[index])
        result['next_cursor'] = token if index + 1 < len(records) else None
        if len(_json(result)) > budget:
            result['items'].pop()
            break
        index += 1
    if not result['items'] and index < len(records):
        _fail('result_limit_exceeded', 'One item exceeds this result budget; increase max_chars.')
    result['next_cursor'] = base64.urlsafe_b64encode(_json([doc['revision'], query, index]).encode()).decode() if index < len(records) else None
    return result


def _inspect(path, sheet, cursor, limit):
    if type(limit) is not int or not 1 <= limit <= 40:
        _fail('invalid_limit', 'limit must be 1..40.')
    doc = _open(path)
    records = [{'kind': 'workbook', 'features': doc['features'], 'reader_warnings': doc['notices'],
                'defined_name_count': len(doc['workbook'].defined_names),
                'calculation': dict(doc['roots']['xl/workbook.xml'].find(_tag('calcPr')).attrib)
                if doc['roots']['xl/workbook.xml'].find(_tag('calcPr')) is not None else {}}]
    if sheet is None:
        for name, part in doc['parts'].items():
            ws, root = _sheet(doc, name), doc['roots'][part]
            records.append({'kind': 'sheet', 'name': name, 'state': ws.sheet_state,
                            'extent': ws.calculate_dimension(), 'stored_cells': len(root.findall('.//' + _tag('c'))),
                            'formula_cells': len(root.findall('.//' + _tag('f'))), 'merged_ranges': len(ws.merged_cells.ranges),
                            'freeze_panes': ws.freeze_panes, 'protected': bool(ws.protection.sheet)})
    else:
        ws = _sheet(doc, sheet)
        records.append({'kind': 'sheet_layout', 'name': sheet, 'state': ws.sheet_state,
                        'default_row_height': ws.sheet_format.defaultRowHeight,
                        'default_column_width': ws.sheet_format.defaultColWidth, 'base_column_width': ws.sheet_format.baseColWidth,
                        'gridlines': ws.sheet_view.showGridLines, 'freeze_panes': ws.freeze_panes,
                        'print_area': str(ws.print_area), 'print_titles': ws.print_titles,
                        'page_setup': dict(ws.page_setup), 'auto_filter': ws.auto_filter.ref})
        for merge in ws.merged_cells.ranges:
            records.append({'kind': 'merge', 'range': str(merge), 'anchor': merge.start_cell.coordinate})
        for index, dim in sorted(ws.row_dimensions.items()):
            records.append({'kind': 'row', 'index': index, 'height': dim.height, 'hidden': dim.hidden, 'outline_level': dim.outlineLevel})
        for key, dim in ws.column_dimensions.items():
            records.append({'kind': 'column', 'key': key, 'min': dim.min, 'max': dim.max,
                            'width': dim.width, 'hidden': dim.hidden, 'outline_level': dim.outlineLevel})
        for scope, names in [('workbook', doc['workbook'].defined_names), ('sheet', ws.defined_names)]:
            for name in names.values():
                records.append({'kind': 'defined_name', 'scope': scope, 'name': name.name, 'definition': _clip(name.attr_text)})
        for table in ws.tables.values():
            records.append({'kind': 'table', 'name': table.name, 'range': table.ref})
        for validation in ws.data_validations.dataValidation:
            records.append({'kind': 'validation', 'range': _clip(str(validation.sqref)), 'type': validation.type,
                            'formula1': _clip(validation.formula1), 'formula2': _clip(validation.formula2)})
        for cf in ws.conditional_formatting:
            records.append({'kind': 'conditional_format', 'range': _clip(str(cf.sqref)),
                            'rule_count': len(ws.conditional_formatting[cf])})
    return _page(doc, records, ['inspect', sheet], cursor, _BUDGET, limit)


def _read(path, revision, sheet, cell_range, cursor, max_chars):
    if type(max_chars) is not int or not 1800 <= max_chars <= _BUDGET:
        _fail('invalid_limit', 'max_chars must be 1800..12000.')
    doc = _open(path, revision)
    ws = _sheet(doc, sheet)
    a, b, c, d = _range(cell_range)
    if (c - a + 1) * (d - b + 1) > 4096:
        _fail('range_limit_exceeded', 'Read at most 4096 cells per rectangle; use narrower ranges.')
    records = []
    for row in ws.iter_rows(min_row=b, max_row=d, min_col=a, max_col=c):
        for cell in row:
            merge = next((m for m in ws.merged_cells.ranges if cell.coordinate in m), None)
            raw = doc['cached'][sheet][cell.coordinate].value if cell.data_type == 'f' else cell.value
            columns = [dim for dim in ws.column_dimensions.values() if dim.min <= cell.column <= dim.max]
            records.append({'cell': cell.coordinate, 'value': _clip(_scalar(raw)),
                            'data_type': cell.data_type, 'formula': _clip(cell.value if type(cell.value) is str else getattr(cell.value, 'text', None)) if cell.data_type == 'f' else None,
                            'formula_kind': type(cell.value).__name__ if cell.data_type == 'f' and type(cell.value) is not str else None,
                            'number_format': cell.number_format, 'style_id': cell.style_id,
                            'merge': str(merge) if merge else None, 'merge_anchor': merge.start_cell.coordinate if merge else None,
                            'row_hidden': ws.row_dimensions[cell.row].hidden if cell.row in ws.row_dimensions else False,
                            'column_hidden': any(dim.hidden for dim in columns),
                            'row_height': ws.row_dimensions[cell.row].height if cell.row in ws.row_dimensions else None,
                            'column_width': columns[0].width if columns else None,
                            'bold': cell.font.bold, 'wrap_text': cell.alignment.wrap_text,
                            'horizontal_alignment': cell.alignment.horizontal,
                            'hyperlink': _clip(cell.hyperlink.target or cell.hyperlink.location) if getattr(cell, 'hyperlink', None) else None,
                            'comment': _clip(cell.comment.text) if getattr(cell, 'comment', None) else None})
    return _page(doc, records, ['read', sheet, cell_range], cursor, max_chars, 4096)


def _admit_derivative(doc):
    if any(doc['features'][k] for k in ('signatures', 'external_links', 'connections', 'query_tables')):
        _fail('unsupported_workbook', 'Signed or externally connected workbooks cannot be edited/rendered.')


def _edit(path, output_path, revision, edits):
    doc = _open(path, revision)
    output = _destination(output_path)
    if output.suffix.lower() != '.xlsx':
        _fail('invalid_path', 'Output must have .xlsx extension.')
    _admit_derivative(doc)
    if any(f.get('t') in ('array', 'dataTable') for part in doc['parts'].values()
           for f in doc['roots'][part].findall('.//' + _tag('f'))):
        _fail('unsupported_workbook', 'Editing books with array/spill/data-table formulas is refused; caches cannot safely be invalidated.')
    if not isinstance(edits, list) or not 1 <= len(edits) <= 64:
        _fail('invalid_edit', 'Supply 1..64 cell edits.')
    seen, changed, receipt = set(), set(), []
    for edit in edits:
        if (type(edit) is not dict or len(edit) != 3 or not {'sheet', 'cell'} <= edit.keys()
                or len({'value', 'formula', 'date'} & edit.keys()) != 1
                or type(edit['sheet']) is not str or type(edit['cell']) is not str):
            _fail('invalid_edit', 'Each edit is {sheet,cell,value}, {sheet,cell,formula} or {sheet,cell,date}.')
        sheet, address = edit['sheet'], edit['cell']
        ws = _sheet(doc, sheet)
        a, b, c, d = _range(address)
        if a != c or b != d or ':' in address:
            _fail('invalid_edit', 'Edit targets one cell, not a range.')
        if (sheet, address) in seen:
            _fail('invalid_edit', 'Duplicate target cell in batch.')
        seen.add((sheet, address))
        if ws.protection.sheet:
            _fail('unsupported_target', 'Editing protected sheets is refused.')
        for merge in ws.merged_cells.ranges:
            if address in merge and address != merge.start_cell.coordinate:
                _fail('merged_cell', 'Edit the merged range anchor only.')
        part = doc['parts'][sheet]
        root = doc['roots'][part]
        data = root.find(_tag('sheetData'))
        cell = data.find(".//" + _tag('c') + "[@r='" + address + "']")
        if cell is not None:
            f = cell.find(_tag('f'))
            if cell.get('cm') or cell.get('vm') or (f is not None and f.attrib):
                _fail('unsupported_target', 'Special/shared formulas and rich cell metadata cannot be edited.')
            if any(e.tag not in {_tag('f'), _tag('v'), _tag('is')} for e in cell):
                _fail('unsupported_target', 'Target has unsupported cell content.')
        else:
            row = data.find(_tag('row') + "[@r='" + str(b) + "']")
            if row is None:
                row = _ET.Element(_tag('row'), r=str(b))
                index = next((i for i, e in enumerate(data) if int(e.get('r')) > b), len(data))
                data.insert(index, row)
            cell = _ET.Element(_tag('c'), r=address)
            index = next((i for i, e in enumerate(row) if _xl.utils.cell.column_index_from_string(_xl.utils.cell.coordinate_from_string(e.get('r'))[0]) > a), len(row))
            row.insert(index, cell)
        field = next(k for k in ('value', 'formula', 'date') if k in edit)
        value = edit[field]
        if field == 'formula':
            if type(value) is not str or not value.startswith('=') or len(value) < 2 or len(value) > 8192 or '[' in value:
                _fail('invalid_edit', 'Formula must start with =, have 2..8192 characters and no external/structured [] references.')
        elif field == 'date':
            try:
                if type(value) is not str:
                    raise ValueError()
                parsed = _dt.datetime.fromisoformat(value) if 'T' in value else _dt.date.fromisoformat(value)
                if isinstance(parsed, _dt.datetime) and parsed.tzinfo is not None:
                    raise ValueError()
            except ValueError:
                _fail('invalid_edit', 'date must be an ISO date or datetime.')
        elif not (value is None or type(value) in (str, bool, int, float)) or (type(value) is float and not math.isfinite(value)):
            _fail('invalid_edit', 'value must be a finite JSON scalar or null.')
        if type(value) in (int, float):
            try:
                if not math.isfinite(float(value)):
                    raise ValueError()
            except (OverflowError, ValueError):
                _fail('invalid_edit', 'Numeric value is outside Excel finite floating-point range.')
        if type(value) is str and len(value) > 32767:
            _fail('invalid_edit', 'Cell text exceeds Excel length limit.')
        for e in list(cell):
            cell.remove(e)
        cell.attrib.pop('t', None)
        try:
            if field == 'formula':
                _ET.SubElement(cell, _tag('f')).text = value[1:]
            elif field == 'date':
                cell.set('t', 'd')
                _ET.SubElement(cell, _tag('v')).text = value
            elif type(value) is str:
                cell.set('t', 'inlineStr')
                text = _ET.SubElement(_ET.SubElement(cell, _tag('is')), _tag('t'))
                text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                text.text = value
            elif value is not None:
                cell.set('t', 'b' if type(value) is bool else 'n')
                _ET.SubElement(cell, _tag('v')).text = str(int(value)) if type(value) is bool else str(value)
        except ValueError:
            _fail('invalid_edit', 'Cell contains invalid XML characters.')
        dimension = root.find(_tag('dimension'))
        if dimension is not None:
            x1, y1, x2, y2 = _range(dimension.get('ref'))
            dimension.set('ref', f'{_xl.utils.get_column_letter(min(x1,a))}{min(y1,b)}:{_xl.utils.get_column_letter(max(x2,a))}{max(y2,b)}')
        changed.add(part)
        receipt.append({'sheet': sheet, 'cell': address, 'kind': field})
    # No dependency engine: invalidate ALL formula caches, including cross-sheet ones.
    for part in doc['parts'].values():
        for cell in doc['roots'][part].findall('.//' + _tag('c')):
            if cell.find(_tag('f')) is not None:
                for v in cell.findall(_tag('v')):
                    cell.remove(v)
                if cell.get('t') in ('str', 'e', 'b'):
                    cell.attrib.pop('t')
                changed.add(part)
    main = doc['roots']['xl/workbook.xml']
    calc = main.find(_tag('calcPr'))
    if calc is None:
        calc = _ET.Element(_tag('calcPr'))
        following = {_tag(n) for n in ('oleSize', 'customWorkbookViews', 'pivotCaches', 'smartTagPr',
                                      'smartTagTypes', 'webPublishing', 'fileRecoveryPr', 'webPublishObjects', 'extLst')}
        index = next((i for i, element in enumerate(main) if element.tag in following), len(main))
        main.insert(index, calc)
    calc.set('fullCalcOnLoad', '1')
    calc.set('forceFullCalc', '1')
    changed.add('xl/workbook.xml')
    removed = {n for n in doc['members'] if n == 'xl/calcChain.xml'}
    for part in ('xl/_rels/workbook.xml.rels', '[Content_Types].xml'):
        for element in list(doc['roots'][part]):
            if element.get('Type', '').endswith('/calcChain') or element.get('PartName') == '/xl/calcChain.xml':
                doc['roots'][part].remove(element)
                changed.add(part)
    updated = {n: v for n, v in doc['members'].items() if n not in removed}
    for part in changed:
        updated[part] = _ET.tostring(doc['roots'][part], xml_declaration=True, encoding='UTF-8', standalone=True)
    buffer = io.BytesIO()
    with ZipFile(buffer, 'w') as archive:
        archive.comment = doc['comment']
        for info in doc['infos']:
            if info.filename not in removed:
                archive.writestr(copy.copy(info), updated[info.filename])
    data = buffer.getvalue()
    if len(data) > 32 * 1024 * 1024:
        _fail('workbook_limit_exceeded', 'Edited copy exceeds 32 MiB.')
    with ZipFile(io.BytesIO(data)) as archive:
        if archive.testzip() is not None or any(archive.read(n) != v for n, v in doc['members'].items() if n not in changed | removed):
            _fail('write_failed', 'Untouched package preservation failed.')
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            checked = _xl.load_workbook(io.BytesIO(data), data_only=False, keep_links=True)
            checked.close()
    except (ValueError, TypeError, KeyError, IndexError) as exc:
        _fail('write_failed', 'Edited workbook failed reader validation: ' + str(exc)[:300])
    result = {'ok': True, 'path': output.relative_to(_ROOT).as_posix(), 'source_revision': doc['revision'],
              'revision': hashlib.sha256(data).hexdigest(), 'edits': receipt, 'changed_parts': sorted(changed),
              'removed_parts': sorted(removed), 'formula_caches': 'invalidated_workbook_wide',
              'untouched_parts_verified': len(doc['members']) - len(changed | removed),
              'warnings': ['Pivot/chart caches and validation rules are preserved, not refreshed or enforced.']}
    if len(_json(result)) > _BUDGET:
        _fail('result_limit_exceeded', 'Edit receipt is too large; use a smaller batch.')
    with tempfile.NamedTemporaryFile(dir=output.parent, prefix='.xlsx-edit-', delete=False) as stream:
        temporary = Path(stream.name)
        try:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
            try:
                os.link(temporary, output)
            except FileExistsError:
                _fail('destination_exists', 'Output appeared during execution.')
        finally:
            temporary.unlink(missing_ok=True)
    return result


def _process(command, stage, deadline):
    with (stage / 'render.log').open('ab') as log:
        proc = subprocess.Popen(command, stdout=log, stderr=log, start_new_session=True)
        try:
            code = proc.wait(timeout=max(.01, deadline - time.monotonic()))
        except subprocess.TimeoutExpired:
            os.killpg(proc.pid, signal.SIGTERM)
            try:
                proc.wait(timeout=.5)
            except subprocess.TimeoutExpired:
                os.killpg(proc.pid, signal.SIGKILL)
                proc.wait()
            _fail('render_timeout', 'Conversion timed out; no output published.')
    if code:
        _fail('render_failed', f'Converter exit {code}: ' + (stage / 'render.log').read_text(errors='replace')[-700:])


def _publish(stage, output):
    libc = ctypes.CDLL(None, use_errno=True)
    rename = getattr(libc, 'renameat2', None)
    if rename is None:
        _fail('dependency_missing', 'Linux renameat2 is required for atomic no-overwrite publication.')
    rename.argtypes = [ctypes.c_int, ctypes.c_char_p, ctypes.c_int, ctypes.c_char_p, ctypes.c_uint]
    rename.restype = ctypes.c_int
    if rename(-100, os.fsencode(stage), -100, os.fsencode(output), 1):
        code = ctypes.get_errno()
        _fail('destination_exists' if code == errno.EEXIST else 'write_failed', os.strerror(code))


def _render(path, revision, output_dir, dpi, timeout_seconds):
    doc = _open(path, revision)
    output = _destination(output_dir)
    _admit_derivative(doc)
    if type(dpi) is not int or not 72 <= dpi <= 200 or type(timeout_seconds) is not int or not 1 <= timeout_seconds <= 300:
        _fail('invalid_render_options', 'Use dpi 72..200 and timeout_seconds 1..300.')
    if not _SOFFICE or not shutil.which('pdftoppm') or not shutil.which('pdfinfo'):
        _fail('dependency_missing', 'LibreOffice soffice and Poppler pdfinfo/pdftoppm must be on PATH before launch.')
    try:
        from pdf2image import convert_from_path as _convert, pdfinfo_from_path as _pdfinfo
        from pdf2image.exceptions import PDFPopplerTimeoutError as _PDFTimeout
        from PIL import Image as _Image
    except ImportError:
        _fail('dependency_missing', 'pdf2image and Pillow are required for rendering.')
    stage = Path(tempfile.mkdtemp(dir=output.parent, prefix='.xlsx-render-'))
    deadline = time.monotonic() + timeout_seconds
    try:
        input_dir, recalculated = stage / 'input', stage / 'calculated'
        input_dir.mkdir()
        recalculated.mkdir()
        (input_dir / 'workbook.xlsx').write_bytes(doc['data'])
        common = [_SOFFICE, '-env:UserInstallation=' + (stage / 'profile').as_uri(), '--headless', '--nologo', '--nodefault', '--norestore']
        _process(common + ['--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', str(recalculated), str(input_dir / 'workbook.xlsx')], stage, deadline)
        calculated = recalculated / 'workbook.xlsx'
        if not calculated.is_file():
            _fail('render_failed', 'LibreOffice did not produce a recalculated workbook.')
        # Inspect the derivative with the same bounded admission before exposing it.
        derivative = _open(calculated.relative_to(_ROOT).as_posix())
        _process(common + ['--convert-to', 'pdf:calc_pdf_Export', '--outdir', str(stage), str(calculated)], stage, deadline)
        pdf = stage / 'workbook.pdf'
        if not pdf.is_file():
            _fail('render_failed', 'LibreOffice did not produce a PDF.')
        try:
            info = _pdfinfo(str(pdf), timeout=max(.01, deadline - time.monotonic()))
            count = info['Pages']
            if not 1 <= count <= 100:
                _fail('render_limit_exceeded', 'Print layout must have 1..100 pages.')
            pages = _convert(str(pdf), dpi=dpi, output_folder=str(stage), fmt='png', paths_only=True,
                             thread_count=1, timeout=max(.01, deadline - time.monotonic()))
            for i, path in enumerate(pages, 1):
                page = Path(path)
                with _Image.open(page) as image:
                    image.verify()
                page.rename(stage / f'page-{i}.png')
        except _PDFTimeout:
            _fail('render_timeout', 'PDF rasterization timed out.')
        except _Error:
            raise
        except Exception as exc:
            _fail('render_failed', 'PDF/PNG verification failed: ' + str(exc)[:500])
        if len(pages) != count:
            _fail('render_failed', 'Incomplete PDF page rasterization.')
        calculated.rename(stage / 'recalculated.xlsx')
        shutil.rmtree(input_dir)
        shutil.rmtree(recalculated)
        shutil.rmtree(stage / 'profile', ignore_errors=True)
        manifest = {'source_revision': doc['revision'], 'recalculated_revision': derivative['revision'],
                    'renderer': 'LibreOffice Calc + Poppler', 'page_count': count, 'dpi': dpi,
                    'pages': [f'page-{i}.png' for i in range(1, count + 1)], 'pdf': 'workbook.pdf',
                    'recalculated_workbook': 'recalculated.xlsx', 'visual_verification': 'not_performed',
                    'print_scope': 'visible sheets according to workbook print settings; no hidden-sheet override',
                    'warnings': ['LibreOffice rewrites the derivative; Excel layout/features/formulas may differ.',
                                 'Recalculated caches are engine results, not certified formula correctness.']}
        (stage / 'manifest.json').write_text(_json(manifest), encoding='utf-8')
        result = {'ok': True, 'revision': doc['revision'], 'recalculated_revision': derivative['revision'],
                  'recalculated_path': (output / 'recalculated.xlsx').relative_to(_ROOT).as_posix(),
                  'manifest_path': (output / 'manifest.json').relative_to(_ROOT).as_posix(),
                  'pdf_path': (output / 'workbook.pdf').relative_to(_ROOT).as_posix(),
                  'page_pattern': (output / 'page-{N}.png').relative_to(_ROOT).as_posix(),
                  'page_count': count, 'visual_verification': 'not_performed', 'warnings': manifest['warnings']}
        if len(_json(result)) > _BUDGET:
            _fail('result_limit_exceeded', 'Render receipt exceeds result limit.')
        _publish(stage, output)
        return result
    finally:
        if stage.exists():
            shutil.rmtree(stage)


def xlsx_inspect(path: str, sheet: str | None = None, cursor: str | None = None, limit: int = 20) -> dict:
    """Map XLSX sheets, extents, hidden/protected state and feature warnings. Select sheet for paginated merges/dimensions/names/print settings; no semantic analysis."""
    return _call(_inspect, path, sheet, cursor, limit)


def xlsx_read(path: str, revision: str, sheet: str, cell_range: str, cursor: str | None = None, max_chars: int = 8000) -> dict:
    """Read an A1 rectangle (<=4096 cells), row-major, with formulas, stored values, number formats, merges and layout hints. Follow cursor with same revision/sheet/range. Caches may be stale; no calculation."""
    return _call(_read, path, revision, sheet, cell_range, cursor, max_chars)


def xlsx_edit(path: str, output_path: str, revision: str, edits: list[dict]) -> dict:
    """Save a NEW XLSX via 1..64 {sheet,cell,value|formula|date} edits. JSON scalar/null value is literal; formula starts =; date is ISO. Preserves styles/other ZIP payloads, invalidates all formula caches; no recalculation or overwrite."""
    return _call(_edit, path, output_path, revision, edits)


def xlsx_render(path: str, revision: str, output_dir: str, dpi: int = 120, timeout_seconds: int = 120) -> dict:
    """Use local LibreOffice/Poppler to create a separate recalculated XLSX, print-layout PDF/PNGs and manifest. Source unchanged. Read derivative with returned revision, inspect page images through existing GUI; render success is not visual approval."""
    return _call(_render, path, revision, output_dir, dpi, timeout_seconds)
